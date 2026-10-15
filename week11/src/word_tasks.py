# agentic-document-automator/src/word_tasks.py
import copy
import os

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from utils import setup_logging

logger = setup_logging(__name__)

TABLE_PLACEHOLDER = "{{TABLE_PLACEHOLDER}}"


class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    @staticmethod
    def _as_placeholder(key):
        """'REPORT_TITLE' -> '{{REPORT_TITLE}}' (keys that already have braces are kept as-is)."""
        key = str(key)
        return key if key.startswith("{{") and key.endswith("}}") else "{{" + key + "}}"

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """
        Replaces a placeholder in a paragraph.
        Word often splits text into several runs (e.g. '{{REPORT_' + 'TITLE}}'),
        so when the placeholder is not inside a single run, the paragraph text is
        rebuilt into the first run (keeping that run's formatting).
        """
        if placeholder not in paragraph.text:
            return False

        replaced = False
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))
                replaced = True

        if not replaced and paragraph.runs:  # placeholder spans multiple runs
            full_text = paragraph.text.replace(placeholder, str(value))
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""
        logger.debug(f"Replaced '{placeholder}' with '{value}'.")
        return True

    def _build_table(self, document, table_data):
        """Creates a formatted table (header row bold + shaded) at the end of the body."""
        cols = len(table_data[0])
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            logger.warning("Style 'Table Grid' not available, using default table style.")

        for r_idx, row_data in enumerate(table_data):
            row = table.add_row()
            for c_idx, cell_data in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(cell_data)
                run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else None
                if r_idx == 0 and run is not None:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    self._shade_cell(cell, "1F4E79")
                if c_idx > 0:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return table

    @staticmethod
    def _shade_cell(cell, hex_color):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tc_pr.append(shd)

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template, filling placeholders and optional table data.
        `data_placeholders` keys may be written as 'REPORT_TITLE' or '{{REPORT_TITLE}}'.
        `table_data` is a list of lists, where the first sub-list is the header row.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        try:
            document = docx.Document(template_path)
            placeholders = {self._as_placeholder(k): v for k, v in data_placeholders.items()}

            # Replace placeholders in paragraphs
            for paragraph in document.paragraphs:
                for placeholder, value in placeholders.items():
                    self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Replace placeholders in existing table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for placeholder, value in placeholders.items():
                                self._replace_placeholder_in_paragraph(paragraph, placeholder, value)

            # Add dynamic table if provided
            if table_data:
                logger.info("Adding dynamic table to the document.")
                table = self._build_table(document, table_data)

                anchor = next((p for p in document.paragraphs if TABLE_PLACEHOLDER in p.text), None)
                if anchor is not None:
                    # move the table to where the placeholder paragraph is, then remove the placeholder
                    anchor._element.addnext(table._tbl)
                    anchor._element.getparent().remove(anchor._element)
                    logger.info("Table inserted at '{{TABLE_PLACEHOLDER}}'.")
                else:
                    logger.warning("Table placeholder '{{TABLE_PLACEHOLDER}}' not found in template. Table appended to end.")

            # warn about anything that was not filled
            leftovers = [p.text for p in document.paragraphs if "{{" in p.text]
            if leftovers:
                logger.warning(f"Unfilled placeholders remain: {leftovers}")

            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True

        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document (paragraphs and tables, in document order)."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None

        try:
            document = docx.Document(docx_path)
            full_text = []
            body = document.element.body
            for child in body.iterchildren():
                tag = child.tag.split('}')[-1]
                if tag == 'p':
                    full_text.append(docx.text.paragraph.Paragraph(child, document).text)
                elif tag == 'tbl':
                    table = docx.table.Table(child, document)
                    for row in table.rows:
                        full_text.append("\t".join(cell.text for cell in row.cells))  # Tab-separated cell values

            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)

        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
